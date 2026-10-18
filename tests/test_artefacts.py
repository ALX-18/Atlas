"""
Sprint Artefacts — agir sur la fenêtre EXACTE (AR1), puis rédiger et produire des documents
Word et PDF (AR2, AR3).

Aucune vraie fenêtre n'est fermée, réduite ou déplacée : les fenêtres sont simulées.
"""

from __future__ import annotations

import asyncio
import json
import pathlib
import sys

import pytest

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

ALEXIS = "fenetre"


def _chat(message, origine=ALEXIS):
    from api.models import ChatRequest
    from api.routes import chat_endpoint

    return asyncio.run(chat_endpoint(ChatRequest(message=message, history=[], origine=origine)))


def _confirmer(r, accepte=True):
    from api.models import ConfirmationResponse
    from api.routes import confirm_endpoint

    assert r.get("type") == "confirmation_required", r
    return asyncio.run(confirm_endpoint(ConfirmationResponse(confirmation_id=r["confirmation_id"], accepted=accepte)))


# --------------------------------------------------------------------------- #
#  AR1 — fermer, réduire, agrandir, ancrer : la fenêtre exacte
# --------------------------------------------------------------------------- #

@pytest.fixture
def deux_rapports(monkeypatch):
    """Deux fenêtres homonymes ; ce qui leur arrive est noté par hwnd, rien n'est fait."""
    from types import SimpleNamespace

    import tools.window_controller as wc
    from core import propositions

    fenetres = [{"hwnd": 101, "title": "AtlasEssai Rapport A", "process": "python.exe"},
                {"hwnd": 202, "title": "AtlasEssai Rapport B", "process": "python.exe"},
                {"hwnd": 303, "title": "Sans titre - Bloc-notes", "process": "Notepad.exe"}]
    faits = []
    monkeypatch.setattr(wc, "_list_visible_windows", lambda: [f for f in fenetres if f["hwnd"] not in
                                                            {h for a, h in faits if a == "fermer"}])
    monkeypatch.setattr(wc, "_agir_sur_hwnd", lambda action, hwnd, **k: faits.append((action, hwnd)) or True)
    monkeypatch.setattr(wc.gw, "getWindowsWithTitle",
                        lambda t: pytest.fail(f"recherche par morceau de titre : {t!r}"))
    monkeypatch.setattr(wc, "time", SimpleNamespace(sleep=lambda s: None))
    propositions.fermer()
    yield faits
    propositions.fermer()


def test_ar1_homonymes_atlas_demande_laquelle_et_ne_ferme_rien(deux_rapports):
    r = _chat("ferme AtlasEssai Rapport")
    m = r["message"]
    assert deux_rapports == [], "jamais la première au hasard"
    assert r["type"] != "confirmation_required"
    assert "laquelle" in m.lower() and "AtlasEssai Rapport A" in m and "AtlasEssai Rapport B" in m, m


def test_ar1_la_confirmation_montre_le_titre_exact_et_le_programme_puis_seule_b_est_fermee(deux_rapports):
    r = _chat("ferme AtlasEssai Rapport B")
    assert "AtlasEssai Rapport B" in (r.get("target") or "") and "python.exe" in (r.get("target") or ""), r
    assert deux_rapports == []
    _confirmer(r)
    assert deux_rapports == [("fermer", 202)], deux_rapports


def test_ar1_un_choix_dans_la_liste_vise_cette_fenetre(deux_rapports):
    _chat("ferme AtlasEssai Rapport")
    r = _chat("la 2")
    assert "AtlasEssai Rapport B" in (r.get("target") or ""), r
    _confirmer(r)
    assert deux_rapports == [("fermer", 202)]


@pytest.mark.parametrize("phrase,action", [("réduis AtlasEssai Rapport A", "reduire"),
                                           ("agrandis AtlasEssai Rapport A", "agrandir")])
def test_ar1_reduire_agrandir_sur_la_fenetre_exacte(deux_rapports, phrase, action):
    _chat(phrase)
    assert deux_rapports == [(action, 101)], deux_rapports


def test_ar1_reduire_des_homonymes_demande_laquelle(deux_rapports):
    r = _chat("réduis AtlasEssai Rapport")
    assert deux_rapports == [] and "laquelle" in r["message"].lower(), r["message"]


def test_ar1_ancrer_sur_la_fenetre_exacte_et_refus_si_homonymes():
    import tools.window_controller as wc

    fenetres = [{"hwnd": 101, "title": "AtlasEssai Rapport A", "process": "python.exe"},
                {"hwnd": 202, "title": "AtlasEssai Rapport B", "process": "python.exe"}]
    faits = []
    import pytest as _p
    mp = _p.MonkeyPatch()
    try:
        mp.setattr(wc, "_list_visible_windows", lambda: fenetres)
        mp.setattr(wc, "_agir_sur_hwnd", lambda action, hwnd, **k: faits.append((action, hwnd)) or True)
        mp.setattr(wc.gw, "getWindowsWithTitle", lambda t: _p.fail("recherche par morceau de titre"))
        r = wc.window_snap("AtlasEssai Rapport", "left")
        assert r["success"] is False and "plusieurs" in r["message"].lower() and faits == []
        r = wc.window_snap("AtlasEssai Rapport B", "left")
        assert r["success"] and faits and faits[0][1] == 202, (r, faits)
    finally:
        mp.undo()


def test_ar1_une_fenetre_qui_a_change_de_titre_depuis_la_confirmation_n_est_pas_fermee(monkeypatch):
    import tools.window_controller as wc

    faits = []
    monkeypatch.setattr(wc, "_list_visible_windows",
                        lambda: [{"hwnd": 202, "title": "Autre chose", "process": "python.exe"}])
    monkeypatch.setattr(wc, "_agir_sur_hwnd", lambda action, hwnd, **k: faits.append((action, hwnd)) or True)
    r = wc.window_close("AtlasEssai Rapport B", hwnd=202)
    assert r["success"] is False and faits == [], r


def test_ar1_title_re_est_echappe_dans_le_grounding(monkeypatch):
    import tools.grounding as g

    vus = []

    class Desktop:
        def __init__(self, backend=None):
            pass

        def windows(self, title_re=None, visible_only=None):
            vus.append(title_re)
            return []

    import pywinauto

    monkeypatch.setattr(pywinauto, "Desktop", Desktop)
    monkeypatch.setattr("tools.window_controller._list_visible_windows", lambda: [])
    asyncio.run(g._try_uia("#général | Atlas (test)", "Bibliothèque"))
    assert vus and all(r is None or ("\\|" in r and "\\(" in r) for r in vus), vus


@pytest.mark.parametrize("nom,processus,attendu", [
    ("steam", ["steamwebhelper.exe"], False),          # un morceau de nom ne suffit plus (GB3)
    ("steam", ["steam.exe"], True),
    ("code", ["vscode-helper.exe"], False),
    ("discord", ["Discord.exe"], True),
    ("Discord.exe", ["discord.exe"], True),
])
def test_ar1_app_lancee_par_nom_de_processus_exact(nom, processus, attendu):
    from core.validator import get_validator

    ctx = {"top_processes": [{"name": p} for p in processus], "foreground_window": {"title": "", "process": ""}}
    assert get_validator()._is_app_running(nom, ctx) is attendu


# --------------------------------------------------------------------------- #
#  AR2 — le moteur déterministe : valider, écrire, relire, ranger
# --------------------------------------------------------------------------- #

CONTENU = {
    "titre": "La photosynthèse",
    "sections": [
        {"titre": "Introduction", "paragraphes": [
            "La photosynthèse est le processus par lequel les plantes transforment la lumière en énergie chimique. "
            "Elle se déroule principalement dans les feuilles, au sein des chloroplastes."]},
        {"titre": "Les deux phases", "paragraphes": [
            "On distingue une phase lumineuse, qui produit de l'énergie, et une phase sombre, qui fixe le carbone."],
         "points": ["Phase lumineuse : capte la lumière", "Phase sombre : cycle de Calvin"]},
        {"titre": "Conclusion", "paragraphes": [
            "Sans la photosynthèse, l'oxygène que nous respirons et la plupart des chaînes alimentaires n'existeraient pas."]},
    ],
}


def test_ar2_un_contenu_structure_est_valide_ou_refuse_avec_la_raison():
    from tools import documents as d

    assert d.valider_contenu(CONTENU) == (True, "")
    for casse, raison in [({**CONTENU, "titre": ""}, "titre"), ({**CONTENU, "sections": []}, "section"),
                          ({**CONTENU, "sections": [{"titre": "x", "paragraphes": []}]}, "paragraphe"),
                          ({**CONTENU, "titre": "Photosynthesis",
                            "sections": [{"titre": "Intro", "paragraphes": ["Photosynthesis is the process by which plants "
                                                                            "turn light into chemical energy in the leaves."]}]},
                           "français"),
                          ("pas un dictionnaire", "structure")]:
        ok, pourquoi = d.valider_contenu(casse)
        assert not ok and raison in pourquoi.lower(), (casse, pourquoi)


def test_ar2_word_vrais_titres_listes_pied_de_page_et_relu(tmp_path):
    import docx

    from tools import documents as d

    f = tmp_path / "essai.docx"
    d.ecrire_docx(CONTENU, f, pied="Rédigé par Atlas avec un modèle local — à relire")
    ok, raison, info = d.verifier_fichier(f)
    assert ok, raison
    doc = docx.Document(str(f))
    styles = [p.style.name for p in doc.paragraphs]
    assert "Title" in styles and styles.count("Heading 1") == 3 and "List Bullet" in styles, styles
    assert "à relire" in doc.sections[0].footer.paragraphs[0].text
    assert "instrText" in doc.element.xml and "TOC" in doc.element.xml          # le sommaire Word


def test_ar2_pdf_accents_titres_et_numeros_de_page(tmp_path):
    from pypdf import PdfReader

    from tools import documents as d

    f = tmp_path / "essai.pdf"
    d.ecrire_pdf(CONTENU, f, pied="Rédigé par Atlas avec un modèle local — à relire")
    ok, raison, info = d.verifier_fichier(f)
    assert ok, raison
    texte = " ".join(p.extract_text() for p in PdfReader(str(f)).pages)
    assert "photosynthèse" in texte and "chaînes" in texte and "à relire" in texte, texte[:300]
    assert "Page 1" in texte


def test_ar2_un_fichier_vide_ou_illisible_n_est_jamais_annonce_cree(tmp_path):
    from tools import documents as d

    vide = tmp_path / "vide.docx"
    vide.write_bytes(b"")
    assert d.verifier_fichier(vide)[0] is False
    faux = tmp_path / "faux.pdf"
    faux.write_bytes(b"pas un pdf")
    assert d.verifier_fichier(faux)[0] is False
    assert d.verifier_fichier(tmp_path / "absent.docx")[0] is False


def test_ar2_rangement_documents_atlas_jamais_d_ecrasement(tmp_path, monkeypatch):
    from tools import documents as d

    monkeypatch.setattr(d, "dossier_atlas", lambda: tmp_path / "Atlas")
    a = d.chemin_libre("La photosynthèse : une introduction / test", ".docx", date="2026-10-17")
    assert a.parent == tmp_path / "Atlas" and a.name == "La photosynthèse une introduction test - 2026-10-17.docx", a
    a.write_bytes(b"x")
    b = d.chemin_libre("La photosynthèse : une introduction / test", ".docx", date="2026-10-17")
    assert b.name.endswith(" (2).docx") and not b.exists()


# --------------------------------------------------------------------------- #
#  AR3 — « rien d'inventé » : chiffres, dates, noms propres présents dans la source
# --------------------------------------------------------------------------- #

SOURCE = ("Le stage s'est déroulé chez Airbus à Toulouse du 3 mars au 28 juin 2026. Encadré par Mme Durand, "
          "le stagiaire a réduit de 15 % le temps de calcul et livré 4 modules. Le budget était de 12 000 euros.")


@pytest.mark.parametrize("resume,inventes", [
    ("Le stage chez Airbus à Toulouse a duré du 3 mars au 28 juin 2026 ; le temps de calcul a baissé de 15 %.", []),
    ("Encadré par Mme Durand, le stagiaire a livré 4 modules pour un budget de 12 000 euros.", []),
    ("Le stage chez Thales a permis de livrer 6 modules.", ["Thales", "6"]),
    ("Le stage a commencé le 12 février 2025 à Paris.", ["12 février 2025", "Paris"]),
])
def test_ar3_faits_absents_de_la_source_detectes(resume, inventes):
    from tools import documents as d

    trouves = d.faits_absents(resume, SOURCE)
    assert sorted(trouves) == sorted(inventes), trouves


def test_ar3_un_resume_avec_un_fait_invente_n_est_pas_livre_tel_quel():
    from tools import documents as d

    contenu = {"titre": "Résumé du stage", "sections": [{"titre": "Bilan", "paragraphes": [
        "Le stage s'est déroulé chez Airbus à Toulouse. Il a permis de livrer 6 modules. Le temps de calcul a baissé de 15 %."]}]}
    propre, retirees = d.retirer_faits_inventes(contenu, SOURCE)
    texte = " ".join(propre["sections"][0]["paragraphes"])
    assert "6 modules" not in texte and "Airbus" in texte and "15 %" in texte, texte
    assert len(retirees) == 1 and "6 modules" in retirees[0]["phrase"] and retirees[0]["faits"] == ["6"]


# --------------------------------------------------------------------------- #
#  AR2 — de la demande au fichier ouvert, sur /api/chat (modèle simulé, ouverture comptée)
# --------------------------------------------------------------------------- #

def _plan(titre="La photosynthèse", n=4):
    return {"titre": titre, "sections": [f"Partie {k}" for k in range(1, n + 1)]}


def _section(k=1):
    return {"paragraphes": [f"Cette partie {k} explique comment les plantes utilisent la lumière du soleil pour "
                            "produire leur énergie, et pourquoi ce mécanisme est essentiel à la vie sur la Terre."],
            "points": ["Les feuilles captent la lumière", "Le carbone est fixé dans les sucres"]}


@pytest.fixture
def atelier(tmp_path, monkeypatch):
    """Dossier Atlas jetable ; le modèle simulé ; le travail fait tout de suite ; l'ouverture comptée."""
    import core.artefacts as ar
    import tools.documents as d
    import tools.fichiers as tf
    from core import annonces, propositions

    monkeypatch.setattr(d, "dossier_atlas", lambda: tmp_path / "Documents" / "Atlas")
    monkeypatch.setattr(ar, "_en_arriere_plan", lambda f: f())
    ouverts = []
    monkeypatch.setattr(tf, "_ouvrir_systeme", lambda c: ouverts.append(c))
    reponses = []

    def modele(systeme, consigne, schema, **k):
        r = reponses.pop(0) if reponses else None
        return r(consigne) if callable(r) else r

    monkeypatch.setattr(ar, "_generer_json", modele)
    annonces.lire()
    propositions.fermer()
    yield {"dossier": tmp_path / "Documents" / "Atlas", "ouverts": ouverts, "reponses": reponses, "ar": ar}
    propositions.fermer()


def _annonces():
    from core import annonces

    return [a["message"] for a in annonces.lire()]


@pytest.mark.parametrize("phrase,type_,sujet,source", [
    ("fais-moi un rapport sur la photosynthèse", "rapport", "la photosynthèse", None),
    ("rédige un rapport en PDF sur l'histoire de Rome", "rapport", "l'histoire de Rome", None),
    ("fais un document Word sur les volcans", "rapport", "les volcans", None),
    ("résume mon rapport de stage dans un Word", "resume", None, "mon rapport de stage"),
    ("fais un résumé en PDF de mon rapport de stage", "resume", None, "mon rapport de stage"),
    ("rédige un compte rendu de mes notes de réunion", "compte_rendu", None, "mes notes de réunion"),
])
def test_ar2_la_demande_est_reconnue(phrase, type_, sujet, source):
    from core.intent_classifier import get_classifier

    i = get_classifier().classify(phrase, {})
    assert i.category == "artefact", (phrase, i.category, i.verb, i.params)
    assert i.params["type"] == type_
    if sujet:
        assert i.params["sujet"].lower() == sujet.lower(), i.params
    if source:
        assert i.params["source"].lower() == source.lower(), i.params


@pytest.mark.parametrize("phrase", ["résume-moi la révolution française", "fais-moi un résumé de la guerre de 14",
                                    "c'est quoi un rapport de force ?", "ouvre mon rapport de stage",
                                    "résume la page https://example.com", "raconte-moi l'histoire de Rome"])
def test_ar2_sans_mot_de_document_ce_n_est_pas_un_artefact(phrase):
    from core.intent_classifier import get_classifier

    assert get_classifier().classify(phrase, {}).category != "artefact", phrase


def test_ar2_un_rapport_annonce_l_attente_puis_cree_verifie_ouvre_et_le_dit(atelier):
    atelier["reponses"] += [_plan(n=4)] + [_section(k) for k in range(1, 5)]
    r = _chat("fais-moi un rapport sur la photosynthèse")
    assert "je rédige" in r["message"].lower() and "secondes" in r["message"].lower(), r["message"]
    (f,) = list(atelier["dossier"].glob("*.docx"))
    assert f.name.startswith("La photosynthèse - ") and atelier["ouverts"] == [str(f)]
    (a,) = _annonces()
    assert a.startswith("J'ai créé") and str(f) in a, a
    import docx

    assert [p.text for p in docx.Document(str(f)).paragraphs if p.style.name == "Heading 1"] == \
        ["Partie 1", "Partie 2", "Partie 3", "Partie 4"]


def test_ar2_pdf_si_demande_les_deux_si_demande(atelier):
    atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
    _chat("rédige un rapport en PDF sur la photosynthèse")
    assert [f.suffix for f in atelier["dossier"].iterdir()] == [".pdf"]
    atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
    _chat("fais un rapport en Word et en PDF sur la photosynthèse")
    assert sorted(f.suffix for f in atelier["dossier"].iterdir()) == [".docx", ".pdf", ".pdf"]


def test_ar2_jamais_d_ecrasement(atelier):
    for _ in range(2):
        atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
        _chat("fais-moi un rapport sur la photosynthèse")
    noms = sorted(f.name for f in atelier["dossier"].glob("*.docx"))
    assert len(noms) == 2 and any(n.endswith(" (2).docx") for n in noms), noms


def test_ar2_contenu_invalide_une_seule_nouvelle_tentative_puis_l_echec_est_dit(atelier):
    atelier["reponses"] += [{"titre": "", "sections": []}, {"pas": "un plan"}]
    _chat("fais-moi un rapport sur la photosynthèse")
    assert list(atelier["dossier"].glob("*")) == [] and atelier["ouverts"] == []
    (a,) = _annonces()
    assert "j'ai créé" not in a.lower() and "pas pu" in a.lower(), a


def test_ar2_un_fichier_qui_ne_se_relit_pas_n_est_jamais_annonce_cree(atelier, monkeypatch):
    import tools.documents as d

    atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
    monkeypatch.setattr(d, "ecrire_docx", lambda contenu, chemin, pied="": pathlib.Path(chemin).write_bytes(b""))
    _chat("fais-moi un rapport sur la photosynthèse")
    (a,) = _annonces()
    assert "j'ai créé" not in a.lower() and atelier["ouverts"] == [], a


def test_ar2_sujet_absent_atlas_demande_avant_de_rediger(atelier):
    r = _chat("fais-moi un rapport")
    assert "quel sujet" in r["message"].lower() and atelier["reponses"] == [] and _annonces() == [], r["message"]
    atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
    r = _chat("la photosynthèse")
    assert "je rédige" in r["message"].lower(), r["message"]
    assert list(atelier["dossier"].glob("*.docx"))


def test_ar2_le_document_cree_entre_dans_l_index(atelier):
    import core.file_indexer as fi

    atelier["reponses"] += [_plan(titre="Les volcans actifs", n=3)] + [_section(k) for k in range(1, 4)]
    _chat("fais un document Word sur les volcans")
    assert any(r["nom"].startswith("Les volcans actifs") for r in fi.chercher_nom("volcans actifs", limite=5))


# --------------------------------------------------------------------------- #
#  AR2.5 / AR3 — résumé d'un fichier de l'index, rien d'inventé
# --------------------------------------------------------------------------- #

@pytest.fixture
def source_indexee(atelier, tmp_path):
    """Un « rapport de stage » de test, indexé."""
    import core.file_indexer as fi

    docs = tmp_path / "pc" / "Documents"
    docs.mkdir(parents=True)
    f = docs / "Rapport de stage AtlasEssai.txt"
    f.write_text(SOURCE + "\n\n" + "Le stagiaire a aussi rédigé une documentation complète. " * 20, encoding="utf-8")
    idx = fi.FileIndexer()
    idx._cfg = {**fi.charger_config(), "scan_roots": [str(tmp_path / "pc")], "exclude_dirs": [], "cpu_max_pourcent": 100}
    fi.oublier_index()
    idx.passage()
    yield f
    fi.oublier_index()


def test_ar3_resume_d_un_fichier_un_fait_invente_est_retire_et_dit(atelier, source_indexee):
    atelier["reponses"] += [
        {"points": ["Stage chez Airbus à Toulouse", "Temps de calcul réduit de 15 %", "6 modules livrés"]},
        {"titre": "Résumé du rapport de stage", "sections": [{"titre": "Bilan du stage", "paragraphes": [
            "Le stage s'est déroulé chez Airbus à Toulouse et a été encadré par Mme Durand. Il a permis de livrer "
            "6 modules. Le temps de calcul a baissé de 15 % et le budget était de 12 000 euros."]}]},
    ]
    r = _chat("résume mon rapport de stage AtlasEssai dans un Word")
    assert "je rédige" in r["message"].lower() or "je résume" in r["message"].lower(), r["message"]
    (f,) = list(atelier["dossier"].glob("*.docx"))
    import docx

    texte = " ".join(p.text for p in docx.Document(str(f)).paragraphs)
    assert "6 modules" not in texte and "Airbus" in texte and "12 000 euros" in texte, texte
    (a,) = _annonces()
    assert "1 phrase" in a and "absent" in a.lower(), a


def test_ar3_resume_d_un_long_document_decoupe_en_morceaux(atelier, source_indexee):
    long = source_indexee.parent / "Rapport long AtlasEssai.txt"
    long.write_text(("Le projet Atlas a avancé chaque semaine. " * 60 + "\n\n") * 12, encoding="utf-8")
    import core.file_indexer as fi
    idx = fi.FileIndexer()
    idx._cfg = {**fi.charger_config(), "scan_roots": [str(source_indexee.parent)], "exclude_dirs": [], "cpu_max_pourcent": 100}
    idx.passage()
    consignes = []

    def modele(consigne):
        if "Rédige un résumé structuré" in consigne:
            return {"titre": "Résumé du rapport long", "sections": [{"titre": "Synthèse", "paragraphes": [
                "Le projet Atlas a avancé chaque semaine, de façon régulière et suivie. Les étapes se sont "
                "enchaînées sans interruption, et chaque semaine a apporté une avancée mesurable du projet."]}]}
        consignes.append(consigne)
        return {"points": ["Le projet Atlas a avancé chaque semaine"]}

    atelier["reponses"] += [modele] * 50
    _chat("résume mon rapport long AtlasEssai dans un Word")
    assert len(consignes) >= 3, "un document plus long que le contexte est découpé"
    assert list(atelier["dossier"].glob("*.docx"))


def test_ar2_source_introuvable_le_dit_sans_rien_rediger(atelier, source_indexee):
    r = _chat("résume mon mémoire de licence dans un Word")
    assert "aucun fichier" in r["message"].lower() and atelier["reponses"] == [] and not list(atelier["dossier"].glob("*"))


def test_ar2_compte_rendu_sans_date_ni_participants_atlas_demande(atelier, source_indexee):
    r = _chat("rédige un compte rendu de mon rapport de stage AtlasEssai")
    assert "date" in r["message"].lower() and "particip" in r["message"].lower(), r["message"]
    assert not list(atelier["dossier"].glob("*"))


def test_ar4_bibliotheque_absente_bascule_dite_et_rien_d_annonce_cree(atelier, monkeypatch):
    import tools.documents as d
    from core_conversational.degradations import actives

    def sans_docx(*a, **k):
        raise ImportError("No module named 'docx'", name="docx")

    monkeypatch.setattr(d, "ecrire_docx", sans_docx)
    atelier["reponses"] += [_plan(n=3)] + [_section(k) for k in range(1, 4)]
    _chat("fais-moi un rapport sur la photosynthèse")
    (a,) = _annonces()
    assert "j'ai créé" not in a.lower() and "manque" in a.lower(), a
    assert "documents" in [b["cle"] for b in actives()]
