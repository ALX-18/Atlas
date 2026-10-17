"""
Sprint Artefacts — le moteur DÉTERMINISTE des documents.

Le modèle de langue rédige un contenu STRUCTURÉ ; ce module, sans modèle :
- le valide (`valider_contenu`) : champs présents, longueurs raisonnables, français ;
- produit le fichier : Word (`ecrire_docx`, python-docx : vrais styles de titre, sommaire Word,
  listes, pied de page numéroté) ou PDF (`ecrire_pdf`, fpdf2 : polices TrueType de Windows pour
  les accents, titres, numéros de page) — sans Word ni LibreOffice installés ;
- le RELIT (`verifier_fichier`) : Atlas ne dit « j'ai créé » qu'après cette relecture ;
- le range (`chemin_libre`) : Documents\\Atlas\\<titre court> - <date>.<ext>, jamais d'écrasement ;
- vérifie qu'un résumé n'invente rien (`faits_absents`, `retirer_faits_inventes`) : nombres,
  dates et noms propres du résumé doivent figurer dans la source.

Structure attendue :
    {"titre": str, "sections": [{"titre": str, "paragraphes": [str, ...], "points": [str, ...]}],
     "points_cles": [str, ...]}         (points, points_cles : facultatifs)
"""

from __future__ import annotations

import copy
import os
import pathlib
import re
import time
import unicodedata
from typing import Any, Optional

PIED_RAPPORT = "Rédigé par Atlas avec un modèle local — à relire"
PIED_RESUME = "Résumé rédigé par Atlas avec un modèle local, vérifié contre le document source — à relire"

_MOTS_FR = {"le", "la", "les", "de", "des", "du", "et", "est", "un", "une", "en", "que", "qui", "pour", "dans",
            "par", "sur", "au", "aux", "ce", "cette", "il", "elle", "sont", "avec", "pas", "plus", "ou", "se",
            "son", "sa", "ses", "leur", "leurs", "on", "nous", "vous", "mais", "comme", "être", "été"}


# --------------------------------------------------------------------------- #
#  Validation du contenu rédigé par le modèle
# --------------------------------------------------------------------------- #

def _textes(contenu: dict) -> list[str]:
    out = [contenu.get("titre", "")]
    for s in contenu.get("sections", []):
        out.append(s.get("titre", ""))
        out += list(s.get("paragraphes", [])) + list(s.get("points", []) or [])
    out += list(contenu.get("points_cles", []) or [])
    return [t for t in out if isinstance(t, str)]


def est_francais(texte: str) -> bool:
    mots = re.findall(r"[a-zàâäéèêëîïôöùûüç']+", texte.lower())
    if len(mots) < 8:
        return True                       # trop court pour juger
    fr = [m for m in mots if m in _MOTS_FR]
    return len(set(fr)) >= 3 and len(fr) / len(mots) >= 0.08


def valider_contenu(contenu: Any) -> tuple[bool, str]:
    """(valide ?, raison). La raison dit ce qui manque : elle est renvoyée au modèle une fois."""
    if not isinstance(contenu, dict):
        return False, "structure : un objet JSON était attendu"
    titre = contenu.get("titre")
    if not isinstance(titre, str) or not (3 <= len(titre.strip()) <= 150):
        return False, "titre absent ou de longueur anormale (3 à 150 caractères)"
    sections = contenu.get("sections")
    if not isinstance(sections, list) or not (1 <= len(sections) <= 15):
        return False, "sections absentes (de 1 à 15 attendues)"
    for k, s in enumerate(sections, 1):
        if not isinstance(s, dict) or not isinstance(s.get("titre"), str) or not s["titre"].strip():
            return False, f"section {k} : titre de section absent"
        paras = s.get("paragraphes")
        if not isinstance(paras, list) or not paras or not all(isinstance(p, str) and 15 <= len(p.strip()) <= 4000
                                                                for p in paras):
            return False, f"section {k} : paragraphes absents ou de longueur anormale (15 à 4000 caractères)"
        points = s.get("points", [])
        if points is not None and (not isinstance(points, list) or not all(isinstance(p, str) and len(p) <= 400
                                                                            for p in points)):
            return False, f"section {k} : liste de points mal formée"
    pc = contenu.get("points_cles", [])
    if pc is not None and (not isinstance(pc, list) or not all(isinstance(p, str) for p in pc)):
        return False, "points clés mal formés"
    tout = " ".join(_textes(contenu))
    from core.ollama_client import contains_non_latin_script

    if contains_non_latin_script(tout) or not est_francais(tout):
        return False, "le texte doit être rédigé en français"
    if len(tout) < 150:
        return False, "contenu trop court"
    return True, ""


# --------------------------------------------------------------------------- #
#  Word
# --------------------------------------------------------------------------- #

def _champ(paragraphe, instruction: str, texte_provisoire: str = "") -> None:
    """Un champ Word (sommaire, numéro de page) : Word le calcule à l'ouverture."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run_avec(el):
        r = OxmlElement("w:r")
        r.append(el)
        paragraphe._p.append(r)

    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    run_avec(debut)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run_avec(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_avec(sep)
    if texte_provisoire:
        t = OxmlElement("w:t")
        t.text = texte_provisoire
        run_avec(t)
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run_avec(fin)


def ecrire_docx(contenu: dict, chemin: pathlib.Path, pied: str = PIED_RAPPORT) -> None:
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = docx.Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    # Langue française (correcteur), et le sommaire mis à jour à l'ouverture.
    rpr = styles["Normal"].element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "fr-FR")
    rpr.append(lang)
    maj = OxmlElement("w:updateFields")
    maj.set(qn("w:val"), "true")
    doc.settings.element.append(maj)
    doc.core_properties.title = contenu["titre"]
    doc.core_properties.author = "Atlas"
    doc.core_properties.language = "fr-FR"

    doc.add_heading(contenu["titre"], level=0)
    if len(contenu["sections"]) >= 3:
        doc.add_paragraph("Sommaire").runs[0].bold = True
        _champ(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
               "Le sommaire se met à jour à l'ouverture (sinon : clic droit, « Mettre à jour les champs »).")
    for s in contenu["sections"]:
        doc.add_heading(s["titre"], level=1)
        for p in s["paragraphes"]:
            doc.add_paragraph(p.strip())
        for pt in s.get("points") or []:
            doc.add_paragraph(pt.strip(), style="List Bullet")
    if contenu.get("points_cles"):
        doc.add_heading("Points clés", level=1)
        for pt in contenu["points_cles"]:
            doc.add_paragraph(pt.strip(), style="List Bullet")
    pied_p = doc.sections[0].footer.paragraphs[0]
    pied_p.text = f"{pied} — page "
    _champ(pied_p, "PAGE", "1")
    chemin.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(chemin))


# --------------------------------------------------------------------------- #
#  PDF
# --------------------------------------------------------------------------- #

_POLICES = {"": "arial.ttf", "B": "arialbd.ttf", "I": "ariali.ttf"}


def ecrire_pdf(contenu: dict, chemin: pathlib.Path, pied: str = PIED_RAPPORT) -> None:
    from fpdf import FPDF

    fonts = pathlib.Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    unicode_ok = all((fonts / f).exists() for f in _POLICES.values())

    def t(texte: str) -> str:            # sans police TrueType : latin-1 seulement (repli)
        return texte if unicode_ok else texte.replace("—", "-").replace("’", "'").replace("œ", "oe") \
            .encode("latin-1", "replace").decode("latin-1")

    class Doc(FPDF):
        def footer(self):
            self.set_y(-14)
            self.set_font(police, "I", 8)
            self.set_text_color(110)
            self.cell(0, 8, t(f"{pied} — Page {self.page_no()}/{{nb}}"), align="C")

    pdf = Doc(format="A4")
    if unicode_ok:
        for style, f in _POLICES.items():
            pdf.add_font("Arial", style, str(fonts / f))
        police = "Arial"
    else:
        police = "Helvetica"
    pdf.set_title(contenu["titre"])
    pdf.set_author("Atlas")
    pdf.set_lang("fr-FR")
    pdf.set_auto_page_break(True, margin=20)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()
    pdf.set_font(police, "B", 20)
    pdf.multi_cell(0, 10, t(contenu["titre"]), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for s in contenu["sections"]:
        pdf.start_section(t(s["titre"]))           # signet (sommaire du lecteur PDF)
        pdf.set_font(police, "B", 14)
        pdf.multi_cell(0, 8, t(s["titre"]), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        pdf.set_font(police, "", 11)
        for p in s["paragraphes"]:
            pdf.multi_cell(0, 6, t(p.strip()), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        for pt in s.get("points") or []:
            pdf.multi_cell(0, 6, t(f"•  {pt.strip()}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    if contenu.get("points_cles"):
        pdf.set_font(police, "B", 14)
        pdf.multi_cell(0, 8, t("Points clés"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(police, "", 11)
        for pt in contenu["points_cles"]:
            pdf.multi_cell(0, 6, t(f"•  {pt.strip()}"), new_x="LMARGIN", new_y="NEXT")
    chemin.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(chemin))


# --------------------------------------------------------------------------- #
#  Relecture : Atlas ne dit « j'ai créé » qu'après elle
# --------------------------------------------------------------------------- #

def verifier_fichier(chemin: pathlib.Path) -> tuple[bool, str, dict]:
    """(bon ?, raison, infos). Le fichier existe, n'est pas vide, se relit et contient du texte."""
    chemin = pathlib.Path(chemin)
    if not chemin.exists():
        return False, "le fichier n'existe pas", {}
    if chemin.stat().st_size == 0:
        return False, "le fichier est vide", {}
    try:
        if chemin.suffix.lower() == ".docx":
            import docx

            doc = docx.Document(str(chemin))
            textes = [p.text for p in doc.paragraphs if p.text.strip()]
            titres = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            if not textes or not titres:
                return False, "le document relu ne contient pas de texte ou de titres", {}
            car = sum(len(x) for x in textes)
            return True, "", {"paragraphes": len(textes), "titres": len(titres), "caracteres": car,
                              "pages_estimees": max(1, round(car / 2600 + 0.5))}
        if chemin.suffix.lower() == ".pdf":
            from pypdf import PdfReader

            r = PdfReader(str(chemin))
            texte = " ".join((p.extract_text() or "") for p in r.pages)
            if not r.pages or len(texte.strip()) < 20:
                return False, "le PDF relu ne contient pas de texte", {}
            return True, "", {"pages": len(r.pages), "caracteres": len(texte)}
    except Exception as e:
        return False, f"le fichier ne se relit pas ({type(e).__name__})", {}
    return False, f"format inattendu : {chemin.suffix}", {}


# --------------------------------------------------------------------------- #
#  Rangement : Documents\Atlas, jamais d'écrasement
# --------------------------------------------------------------------------- #

def dossier_documents() -> pathlib.Path:
    """Le dossier Documents réel (il peut être déplacé, par exemple dans OneDrive)."""
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders") as k:
            v = os.path.expandvars(winreg.QueryValueEx(k, "Personal")[0])
            if v and os.path.isdir(v):
                return pathlib.Path(v)
    except Exception:
        pass
    return pathlib.Path(os.path.expandvars("%USERPROFILE%")) / "Documents"


def dossier_atlas() -> pathlib.Path:
    return dossier_documents() / "Atlas"


def titre_court(titre: str, longueur: int = 60) -> str:
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", titre or "Document")
    t = " ".join(t.split()).strip(" .")
    if len(t) > longueur:
        t = t[:longueur].rsplit(" ", 1)[0]
    return t or "Document"


def chemin_libre(titre: str, ext: str, date: Optional[str] = None) -> pathlib.Path:
    dossier = dossier_atlas()
    dossier.mkdir(parents=True, exist_ok=True)
    base = f"{titre_court(titre)} - {date or time.strftime('%Y-%m-%d')}"
    chemin = dossier / f"{base}{ext}"
    n = 2
    while chemin.exists():
        chemin = dossier / f"{base} ({n}){ext}"
        n += 1
    return chemin


# --------------------------------------------------------------------------- #
#  Rien d'inventé : chiffres, dates, noms propres du résumé présents dans la source
# --------------------------------------------------------------------------- #

_MOIS = "janvier|fevrier|février|mars|avril|mai|juin|juillet|aout|août|septembre|octobre|novembre|decembre|décembre"
_DATE = re.compile(rf"\b(?:1er|\d{{1,2}})\s+(?:{_MOIS})(?:\s+\d{{4}})?\b", re.IGNORECASE)
_NOMBRE = re.compile(r"(?<![\w])\d+(?:[ \u00a0\u202f]\d{3})*(?:[.,]\d+)?(?![\w])")
_PAS_DES_NOMS = {"mme", "mlle", "m", "mr", "dr", "pr", "le", "la", "les", "un", "une", "il", "elle", "ils", "elles",
                 "ce", "cette", "ces", "en", "au", "aux", "du", "de", "des", "et", "or", "mais", "puis", "ainsi",
                 "enfin", "ensuite", "lors", "selon", "pour", "par", "sur", "dans", "avec", "sans", "chez", "on",
                 "nous", "vous", "je", "tu", "son", "sa", "ses", "leur", "leurs", "notre", "votre", "points", "point"}


def _plat(texte: str) -> str:
    t = unicodedata.normalize("NFD", (texte or "").lower())
    return " ".join("".join(c for c in t if unicodedata.category(c) != "Mn").split())


def _nombre_norm(n: str) -> str:
    n = re.sub(r"[ \u00a0\u202f]", "", n).replace(",", ".")
    return n[:-2] if n.endswith(".0") else n


def _noms_propres(texte: str) -> list[str]:
    """Mots à majuscule qui ne commencent pas une phrase (« Airbus », « Toulouse », « Durand »)."""
    out = []
    for phrase in re.split(r"(?<=[.!?;:])\s+|\n+", texte):
        mots = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ][\w'’-]*", phrase)
        for k, m in enumerate(mots):
            if k == 0:
                continue                   # début de phrase : majuscule de grammaire
            if m[0].isupper() and m.lower().strip("'’") not in _PAS_DES_NOMS and len(m) > 1:
                out.append(m.strip("'’"))
    return out


def faits_absents(resume: str, source: str) -> list[str]:
    """Les faits du résumé (dates, nombres, noms propres) qui ne sont PAS dans la source."""
    src = _plat(source)
    nombres_src = {_nombre_norm(n) for n in _NOMBRE.findall(source)}
    absents = []
    reste = resume
    for d in _DATE.findall(resume):
        if _plat(d) not in src:
            absents.append(d)
        reste = reste.replace(d, " ")
    for n in _NOMBRE.findall(reste):
        if _nombre_norm(n) not in nombres_src:
            absents.append(n.strip())
    mots_src = set(re.findall(r"[\w'’-]+", src))
    for nom in _noms_propres(reste):
        if _plat(nom) not in mots_src and _plat(nom) not in src:
            absents.append(nom)
    return list(dict.fromkeys(absents))


def _phrases(texte: str) -> list[str]:
    return [p for p in re.split(r"(?<=[.!?])\s+", texte.strip()) if p]


def retirer_faits_inventes(contenu: dict, source: str) -> tuple[dict, list[dict]]:
    """Retire du résumé chaque PHRASE qui cite un fait absent de la source. Retourne le contenu
    nettoyé et la liste des phrases retirées (dites à Alexis). Un paragraphe vidé disparaît ;
    un point inventé aussi. Le titre inventé devient « Résumé »."""
    propre = copy.deepcopy(contenu)
    retirees = []
    if faits_absents(propre.get("titre", ""), source):
        propre["titre"] = "Résumé"
    for s in propre.get("sections", []):
        paras = []
        for p in s.get("paragraphes", []):
            gardees = []
            for ph in _phrases(p):
                f = faits_absents(ph, source)
                if f:
                    retirees.append({"phrase": ph, "faits": f})
                else:
                    gardees.append(ph)
            if gardees:
                paras.append(" ".join(gardees))
        s["paragraphes"] = paras
        pts = []
        for pt in s.get("points") or []:
            f = faits_absents(pt, source)
            if f:
                retirees.append({"phrase": pt, "faits": f})
            else:
                pts.append(pt)
        if "points" in s:
            s["points"] = pts
    propre["sections"] = [s for s in propre.get("sections", []) if s.get("paragraphes") or s.get("points")]
    pc = []
    for pt in propre.get("points_cles") or []:
        f = faits_absents(pt, source)
        if f:
            retirees.append({"phrase": pt, "faits": f})
        else:
            pc.append(pt)
    if "points_cles" in propre:
        propre["points_cles"] = pc
    return propre, retirees
